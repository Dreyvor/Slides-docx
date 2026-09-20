import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import cv2
import numpy
from docx import Document
from docx.enum.section import WD_ORIENT

from slides_docx.profiles import ProfileStore, make_profile


@unittest.skipUnless(shutil.which("ffmpeg") and shutil.which("ffprobe"), "FFmpeg is required")
class PipelineIntegrationTests(unittest.TestCase):
    def test_detect_and_build_use_same_crop(self):
        repository = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory(prefix="slides docx integration ") as directory:
            root = Path(directory)
            video = root / "lécture recording.mkv"
            vtt = root / "lécture recording.vtt"
            config_root = root / "config"
            filter_graph = (
                "color=black:s=640x360:r=10:d=1,drawbox=x=480:y=0:w=160:h=360:color=blue:t=fill[a];"
                "color=black:s=640x360:r=10:d=1,drawbox=x=480:y=0:w=160:h=360:color=white:t=fill[b];"
                "color=white:s=640x360:r=10:d=1[c];"
                "color=white:s=640x360:r=10:d=1,drawbox=x=480:y=0:w=160:h=360:color=blue:t=fill[d];"
                "[a][b][c][d]concat=n=4:v=1:a=0"
            )
            subprocess.run(
                ["ffmpeg", "-v", "error", "-f", "lavfi", "-i", filter_graph,
                 "-c:v", "ffv1", str(video)],
                check=True,
            )
            vtt.write_text(
                "WEBVTT\n\n00:00:00.100 --> 00:00:01.500\nFirst slide.\n\n"
                "00:00:02.100 --> 00:00:03.500\nSecond slide.\n",
                encoding="utf-8",
            )
            env = dict(os.environ, PYTHONPATH=str(repository), XDG_CONFIG_HOME=str(config_root))
            store = ProfileStore(config_root / "slides-docx" / "config.json")
            store.set_profile("lecture", make_profile((480, 360, 0, 0), 640, 360))

            detect = subprocess.run(
                [sys.executable, "-m", "slides_docx", "detect", str(video), "--threshold", "8"],
                cwd=repository, env=env, text=True, capture_output=True,
            )
            self.assertEqual(detect.returncode, 0, detect.stdout + detect.stderr)
            times_file = root / "lécture recording.slide-times.txt"
            numeric_lines = [
                float(line) for line in times_file.read_text().splitlines()
                if line and not line.startswith("#")
            ]
            self.assertEqual(len(numeric_lines), 1)
            self.assertAlmostEqual(numeric_lines[0], 2.0, delta=0.15)
            contact_sheet = root / "lécture recording.contact-sheet.jpg"
            self.assertTrue(contact_sheet.is_file())
            contact_image = cv2.imdecode(
                numpy.frombuffer(contact_sheet.read_bytes(), dtype=numpy.uint8),
                cv2.IMREAD_COLOR,
            )
            self.assertEqual(contact_image.shape[:2], (185, 1040))

            contact_sheet.unlink()
            without_contact = subprocess.run(
                [sys.executable, "-m", "slides_docx", "detect", str(video),
                 "--profile", "lecture", "--threshold", "8", "--no-contact-sheet"],
                cwd=repository, env=env, text=True, capture_output=True,
            )
            self.assertEqual(
                without_contact.returncode, 0,
                without_contact.stdout + without_contact.stderr,
            )
            self.assertFalse(contact_sheet.exists())
            saved_profile = store.load()["profiles"]["lecture"]
            self.assertEqual(
                saved_profile["settings"]["detect"],
                {"threshold": 8.0, "contact_sheet": False},
            )

            build = subprocess.run(
                [sys.executable, "-m", "slides_docx", "build", str(video), str(vtt)],
                cwd=repository, env=env, text=True, capture_output=True,
            )
            self.assertEqual(build.returncode, 0, build.stdout + build.stderr)
            output = root / "lécture recording.docx"
            document = Document(output)
            self.assertEqual(len(document.inline_shapes), 2)
            self.assertEqual(
                [section.orientation for section in document.sections],
                [WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT] * 2,
            )
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("First slide.", text)
            self.assertIn("Second slide.", text)
            with zipfile.ZipFile(output) as archive:
                media = [name for name in archive.namelist() if name.startswith("word/media/")]
                self.assertEqual(len(media), 2)
                self.assertTrue(all(name.endswith(".png") for name in media))
                for name in media:
                    image = cv2.imdecode(
                        numpy.frombuffer(archive.read(name), dtype=numpy.uint8),
                        cv2.IMREAD_COLOR,
                    )
                    self.assertEqual(image.shape[:2], (360, 480))

            transcript_only = subprocess.run(
                [
                    sys.executable,
                    "-m",
                    "slides_docx",
                    "build",
                    str(video),
                    str(vtt),
                    "--no-slide-images",
                    "--output",
                    str(root / "official slides notes"),
                ],
                cwd=repository,
                env=env,
                text=True,
                capture_output=True,
            )
            self.assertEqual(
                transcript_only.returncode,
                0,
                transcript_only.stdout + transcript_only.stderr,
            )
            transcript_output = root / "official slides notes.docx"
            transcript_document = Document(transcript_output)
            self.assertEqual(len(transcript_document.inline_shapes), 0)
            self.assertEqual(
                [section.orientation for section in transcript_document.sections],
                [WD_ORIENT.PORTRAIT] * 2,
            )
            transcript_text = "\n".join(
                paragraph.text for paragraph in transcript_document.paragraphs
            )
            self.assertIn("First slide.", transcript_text)
            self.assertIn("Second slide.", transcript_text)
            with zipfile.ZipFile(transcript_output) as archive:
                self.assertFalse(
                    any(name.startswith("word/media/") for name in archive.namelist())
                )

            store.set_profile("lecture", make_profile((400, 360, 0, 0), 640, 360))
            changed = subprocess.run(
                [sys.executable, "-m", "slides_docx", "build", str(video), str(vtt)],
                cwd=repository, env=env, text=True, capture_output=True,
            )
            self.assertEqual(changed.returncode, 2)
            self.assertIn("changed after slide detection", changed.stderr)


if __name__ == "__main__":
    unittest.main()
