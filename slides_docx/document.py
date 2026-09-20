import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.exceptions import UnrecognizedImageError
from docx.shared import Cm, Mm

from .content import read_vtt, screenshot_time, seconds_to_timestamp, transcript_by_slide
from .errors import SlidesDocxError
from .video import extract_frame


def set_landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = section.bottom_margin = Mm(10)
    section.left_margin = section.right_margin = Mm(10)


def set_portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(20)


def add_slide_image(document, image_path):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        paragraph.add_run().add_picture(str(image_path), width=Cm(25))
    except (OSError, UnrecognizedImageError) as exc:
        raise SlidesDocxError(
            f"Could not add the extracted slide image to the DOCX: {image_path}"
        ) from exc


def add_transcript(document, number, start_time, text):
    document.add_heading(
        f"Slide {number:02d} — {seconds_to_timestamp(start_time)}", level=1
    )
    document.add_paragraph(text or "[No transcript assigned to this slide]")


def build_document(
    video,
    vtt,
    slide_times,
    duration,
    output,
    lead=5.0,
    crop=None,
    progress=None,
    cancel=None,
):
    video, vtt, output = Path(video), Path(vtt), Path(output)
    if not vtt.is_file():
        raise SlidesDocxError(f"VTT transcript does not exist: {vtt}")
    if not output.parent.is_dir():
        raise SlidesDocxError(f"Output directory does not exist: {output.parent}")
    if lead < 0:
        raise SlidesDocxError("Screenshot lead time cannot be negative.")
    starts = [0.0] + sorted(set(slide_times))
    transcripts = transcript_by_slide(read_vtt(vtt), starts)
    ranges = []
    for index, start in enumerate(starts):
        end = starts[index + 1] if index + 1 < len(starts) else duration
        ranges.append((start, end, screenshot_time(start, end, lead)))

    document = Document()
    set_landscape(document.sections[0])
    with tempfile.TemporaryDirectory(prefix="lecture_slides_") as directory:
        directory = Path(directory)
        for index, (start, _end, capture) in enumerate(ranges):
            if cancel is not None:
                cancel.raise_if_cancelled()
            number = index + 1
            # PNG avoids python-docx rejecting valid FFmpeg JPEGs that do not
            # contain the narrower JFIF/Exif marker layout it expects.
            screenshot = directory / f"slide_{number:03d}.png"
            message = (
                f"Slide {number:02d}: starts {seconds_to_timestamp(start)}, "
                f"screenshot {seconds_to_timestamp(capture)}"
            )
            if progress:
                progress(number, len(ranges), message)
            extract_frame(video, capture, screenshot, crop=crop, cancel=cancel)
            if number > 1:
                set_landscape(document.add_section(WD_SECTION.NEW_PAGE))
            add_slide_image(document, screenshot)
            set_portrait(document.add_section(WD_SECTION.NEW_PAGE))
            add_transcript(document, number, start, " ".join(transcripts[index]))
    try:
        document.save(output)
    except (OSError, PermissionError) as exc:
        raise SlidesDocxError(f"Could not save DOCX to {output}: {exc}") from exc
    return len(starts)
