#!/usr/bin/env python3

import argparse
import bisect
import html
import re
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Mm


# ------------------------------------------------------------
# Time helpers
# ------------------------------------------------------------

def timestamp_to_seconds(timestamp):
    """
    Convert WebVTT timestamp to seconds.
    Accepts:
        HH:MM:SS.mmm
        MM:SS.mmm
    """

    parts = timestamp.strip().split(":")

    if len(parts) == 3:
        hours = int(parts[0])
        minutes = int(parts[1])
        seconds = float(parts[2])

    elif len(parts) == 2:
        hours = 0
        minutes = int(parts[0])
        seconds = float(parts[1])

    else:
        raise ValueError(f"Invalid timestamp: {timestamp}")

    return hours * 3600 + minutes * 60 + seconds


def seconds_to_timestamp(seconds):
    """
    Convert seconds to HH:MM:SS.mmm
    """

    seconds = max(0, seconds)

    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = seconds % 60

    return f"{hours:02d}:{minutes:02d}:{secs:06.3f}"


# ------------------------------------------------------------
# Video helpers
# ------------------------------------------------------------

def get_video_duration(video):
    """
    Get duration of video using ffprobe.
    """

    command = [
        "ffprobe",
        "-v", "error",
        "-show_entries", "format=duration",
        "-of", "default=noprint_wrappers=1:nokey=1",
        str(video)
    ]

    try:
        result = subprocess.check_output(
            command,
            text=True
        ).strip()

        return float(result)

    except (subprocess.CalledProcessError, ValueError):
        sys.exit("Could not determine video duration with ffprobe.")


def extract_screenshot(video, timestamp, output, crop=None):
    """
    Extract one JPEG frame from the video at timestamp.
    """

    command = [
        "ffmpeg",
        "-hide_banner",
        "-loglevel", "error",
        "-y",

        # Seek before opening input = much faster for long recordings.
        "-ss", f"{timestamp:.3f}",

        "-i", str(video),
    ]

    if crop:
        command += [
            "-vf",
            f"crop={crop}"
        ]

    command += [
        "-frames:v", "1",

        # High-quality JPEG
        "-q:v", "2",

        str(output)
    ]

    subprocess.run(
        command,
        check=True
    )


# ------------------------------------------------------------
# VTT
# ------------------------------------------------------------

def read_vtt(vtt_file):
    """
    Return:
        [
            (start_seconds, text),
            ...
        ]
    """

    with open(vtt_file, encoding="utf-8") as f:
        contents = f.read()

    blocks = re.split(r"\n\s*\n", contents)

    cues = []

    for block in blocks:

        lines = block.strip().splitlines()

        timestamp_index = None

        for i, line in enumerate(lines):
            if "-->" in line:
                timestamp_index = i
                break

        if timestamp_index is None:
            continue

        timestamp_line = lines[timestamp_index]

        start_timestamp = (
            timestamp_line
            .split("-->")[0]
            .strip()
        )

        try:
            start_seconds = timestamp_to_seconds(
                start_timestamp
            )
        except ValueError:
            continue

        text_lines = lines[timestamp_index + 1:]

        text = " ".join(
            line.strip()
            for line in text_lines
            if line.strip()
        )

        # Remove possible WebVTT formatting tags.
        text = re.sub(r"<[^>]+>", "", text)

        text = html.unescape(text)

        # Normalize whitespace.
        text = re.sub(r"\s+", " ", text).strip()

        if text:
            cues.append(
                (start_seconds, text)
            )

    return cues


# ------------------------------------------------------------
# Slide timestamps
# ------------------------------------------------------------

def read_slide_times(filename, video_duration):
    """
    slide_times.txt contains the times at which the NEW slide appears.

    Example:

        35.240
        92.560
        181.400

    Slide 1 therefore implicitly starts at 0.
    """

    times = []

    with open(filename) as f:

        for line in f:

            line = line.strip()

            if not line:
                continue

            try:
                value = float(line)
            except ValueError:
                continue

            # Ignore invalid/out-of-range detections.
            if 0 < value < video_duration:
                times.append(value)

    # Remove duplicates and sort.
    times = sorted(set(times))

    # Slide 1 starts at zero.
    return [0.0] + times


# ------------------------------------------------------------
# Decide which video frame represents each slide
# ------------------------------------------------------------

def screenshot_time(slide_start, slide_end, lead):
    """
    Normally:
        screenshot = slide_end - lead

    If that is too close to the start of a short slide,
    capture about 75% of the way through instead.
    """

    duration = slide_end - slide_start

    if duration <= 0:
        return slide_start

    desired = slide_end - lead

    # Avoid screenshots immediately after a slide appears.
    minimum_safe_time = slide_start + min(
        1.0,
        duration * 0.25
    )

    if desired <= minimum_safe_time:

        # Short slide:
        # take a frame near the end instead.
        desired = slide_start + duration * 0.75

    # Avoid hitting the transition frame itself.
    desired = min(
        desired,
        slide_end - 0.15
    )

    return max(
        slide_start,
        desired
    )


# ------------------------------------------------------------
# Map transcript cues to slides
# ------------------------------------------------------------

def transcript_by_slide(cues, slide_starts):
    """
    Returns one list of transcript strings per slide.
    """

    transcripts = [
        []
        for _ in slide_starts
    ]

    for cue_time, text in cues:

        slide_index = (
            bisect.bisect_right(
                slide_starts,
                cue_time
            )
            - 1
        )

        if 0 <= slide_index < len(transcripts):
            transcripts[slide_index].append(text)

    return transcripts


# ------------------------------------------------------------
# Word formatting
# ------------------------------------------------------------

def set_landscape(section):
    """
    A4 landscape page.
    """

    section.orientation = WD_ORIENT.LANDSCAPE

    # Explicitly swap width/height.
    section.page_width = Mm(297)
    section.page_height = Mm(210)

    section.top_margin = Mm(10)
    section.bottom_margin = Mm(10)
    section.left_margin = Mm(10)
    section.right_margin = Mm(10)


def set_portrait(section):
    """
    A4 portrait page.
    """

    section.orientation = WD_ORIENT.PORTRAIT

    section.page_width = Mm(210)
    section.page_height = Mm(297)

    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)


def add_slide_image(doc, image_path):
    """
    Add image centered on the landscape page.

    25 cm width fits both 16:9 and 4:3 slides
    inside an A4 landscape page with 1 cm margins.
    """

    paragraph = doc.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()

    run.add_picture(
        str(image_path),
        width=Cm(25)
    )


def add_transcript(doc, number, start_time, text):
    """
    Add selectable/editable transcript text.
    """

    heading = doc.add_heading(
        f"Slide {number:02d} — "
        f"{seconds_to_timestamp(start_time)}",
        level=1
    )

    if text:

        paragraph = doc.add_paragraph(text)

    else:

        paragraph = doc.add_paragraph(
            "[No transcript assigned to this slide]"
        )


# ------------------------------------------------------------
# Main
# ------------------------------------------------------------

def parse_course_date(value):
    if not re.fullmatch(r"[0-9]{2}\.[0-9]{2}\.[0-9]{4}", value):
        raise argparse.ArgumentTypeError("Date must use DD.MM.YYYY format")
    try:
        return datetime.strptime(value, "%d.%m.%Y").date()
    except ValueError:
        raise argparse.ArgumentTypeError("Date must be a valid calendar date in DD.MM.YYYY format") from None


def main():

    parser = argparse.ArgumentParser(
        description=(
            "Create a Word document containing "
            "PowerPoint screenshots and VTT transcript."
        )
    )

    parser.add_argument(
        "vtt",
        help="Transcript in WebVTT format"
    )

    parser.add_argument(
        "slide_times",
        help="Text file containing slide-change times in seconds"
    )

    parser.add_argument(
        "video",
        help="Lecture video"
    )

    parser.add_argument(
        "output",
        nargs="?",
        help="Output .docx file (default: video path with its extension replaced by .docx)"
    )

    parser.add_argument(
        "--date",
        type=parse_course_date,
        metavar="DD.MM.YYYY",
        help="Course date; prefix the output filename with YYYY_MM_DD-"
    )

    parser.add_argument(
        "--lead",
        type=float,
        default=5.0,
        help=(
            "Take screenshot this many seconds "
            "before next slide change "
            "(default: 5)"
        )
    )

    parser.add_argument(
        "--crop",
        default=None,
        help=(
            "Optional ffmpeg crop specification "
            "W:H:X:Y, e.g. 1920:1080:0:0"
        )
    )

    args = parser.parse_args()

    vtt_file = Path(args.vtt)
    slide_file = Path(args.slide_times)
    video_file = Path(args.video)
    output_file = Path(args.output) if args.output else video_file.with_suffix(".docx")
    if args.date:
        prefix = f"{args.date.year:04d}_{args.date.month:02d}_{args.date.day:02d}-"
        output_file = output_file.with_name(prefix + output_file.name)

    # ----------------------------------------
    # Read source data
    # ----------------------------------------

    print("Reading video...")

    video_duration = get_video_duration(
        video_file
    )

    print(
        "Video duration:",
        seconds_to_timestamp(video_duration)
    )

    slide_starts = read_slide_times(
        slide_file,
        video_duration
    )

    cues = read_vtt(
        vtt_file
    )

    transcripts = transcript_by_slide(
        cues,
        slide_starts
    )

    print(
        f"Detected {len(slide_starts)} slides"
    )

    # ----------------------------------------
    # Build slide ranges
    # ----------------------------------------

    slide_ranges = []

    for i, start in enumerate(slide_starts):

        if i + 1 < len(slide_starts):
            end = slide_starts[i + 1]
        else:
            end = video_duration

        capture = screenshot_time(
            start,
            end,
            args.lead
        )

        slide_ranges.append(
            (start, end, capture)
        )

    # ----------------------------------------
    # Generate Word document
    # ----------------------------------------

    document = Document()

    # First page is slide 1 landscape.
    set_landscape(
        document.sections[0]
    )

    with tempfile.TemporaryDirectory(
        prefix="lecture_slides_"
    ) as tmpdir:

        tmpdir = Path(tmpdir)

        for i, (
            start,
            end,
            capture
        ) in enumerate(slide_ranges):

            slide_number = i + 1

            screenshot = (
                tmpdir
                / f"slide_{slide_number:03d}.jpg"
            )

            print(
                f"Slide {slide_number:02d}: "
                f"starts {seconds_to_timestamp(start)}, "
                f"screenshot {seconds_to_timestamp(capture)}"
            )

            extract_screenshot(
                video_file,
                capture,
                screenshot,
                crop=args.crop
            )

            # --------------------------------
            # Landscape screenshot page
            # --------------------------------

            if slide_number > 1:

                landscape_section = (
                    document.add_section(
                        WD_SECTION.NEW_PAGE
                    )
                )

                set_landscape(
                    landscape_section
                )

            add_slide_image(
                document,
                screenshot
            )

            # --------------------------------
            # Portrait transcript page
            # --------------------------------

            portrait_section = (
                document.add_section(
                    WD_SECTION.NEW_PAGE
                )
            )

            set_portrait(
                portrait_section
            )

            text = " ".join(
                transcripts[i]
            )

            add_transcript(
                document,
                slide_number,
                start,
                text
            )

    # ----------------------------------------
    # Save
    # ----------------------------------------

    document.save(
        output_file
    )

    print()
    print(
        f"Created: {output_file}"
    )


if __name__ == "__main__":
    main()
